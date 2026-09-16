import io
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

from utils.metrics_catalog import SECTION_A, SECTION_B, SECTION_C, PRINCIPLES

HEADER_FILL = PatternFill(start_color="1F3864", end_color="1F3864", fill_type="solid")
HEADER_FONT = Font(color="FFFFFF", bold=True, size=11)
TITLE_FONT = Font(bold=True, size=14, color="1F3864")
SECTION_FONT = Font(bold=True, size=12, color="FFFFFF")
SECTION_FILL = PatternFill(start_color="2E75B6", end_color="2E75B6", fill_type="solid")


def _style_header(ws, row, col_count):
    for c in range(1, col_count + 1):
        cell = ws.cell(row=row, column=c)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center")


def _section_banner(ws, row, text, col_span=3):
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=col_span)
    cell = ws.cell(row=row, column=1, value=text)
    cell.font = SECTION_FONT
    cell.fill = SECTION_FILL
    cell.alignment = Alignment(horizontal="left", vertical="center")


def build_excel_report(org, period, esg_rows):
    """
    esg_rows: list of ESGData model instances for the org+period.
    Returns a BytesIO buffer containing the .xlsx workbook.
    """
    values_by_code = {row.metric_code: row for row in esg_rows}

    wb = Workbook()
    ws = wb.active
    ws.title = "BRSR Report"
    ws.column_dimensions["A"].width = 55
    ws.column_dimensions["B"].width = 20
    ws.column_dimensions["C"].width = 15

    ws.merge_cells("A1:C1")
    ws["A1"] = f"BRSR Disclosure Report — {org.name}"
    ws["A1"].font = TITLE_FONT
    ws.merge_cells("A2:C2")
    ws["A2"] = f"Reporting Period: {period.name}  |  Org Type: {org.org_type.title()}  |  Code: {org.code}"

    row = 4
    _section_banner(ws, row, "SECTION A: General Disclosures")
    row += 1
    ws.cell(row=row, column=1, value="Field")
    ws.cell(row=row, column=2, value="Value")
    ws.cell(row=row, column=3, value="Unit")
    _style_header(ws, row, 3)
    row += 1
    for m in SECTION_A:
        entry = values_by_code.get(m["code"])
        ws.cell(row=row, column=1, value=m["name"])
        ws.cell(row=row, column=2, value=entry.value if entry else "")
        ws.cell(row=row, column=3, value=m.get("unit", ""))
        row += 1

    row += 1
    _section_banner(ws, row, "SECTION B: Management & Process Disclosures (Policy Coverage)")
    row += 1
    ws.cell(row=row, column=1, value="Principle Policy Question")
    ws.cell(row=row, column=2, value="Response")
    ws.cell(row=row, column=3, value="")
    _style_header(ws, row, 3)
    row += 1
    for m in SECTION_B:
        entry = values_by_code.get(m["code"])
        ws.cell(row=row, column=1, value=m["name"])
        ws.cell(row=row, column=2, value=entry.value if entry else "")
        row += 1

    row += 1
    _section_banner(ws, row, "SECTION C: Principle-wise Performance Metrics")
    row += 1
    ws.cell(row=row, column=1, value="Metric")
    ws.cell(row=row, column=2, value="Value")
    ws.cell(row=row, column=3, value="Unit")
    _style_header(ws, row, 3)
    row += 1
    current_principle = None
    for m in SECTION_C:
        if m["principle"] != current_principle:
            current_principle = m["principle"]
            ws.cell(row=row, column=1, value=f"{current_principle}: {PRINCIPLES[current_principle]}").font = Font(bold=True, italic=True)
            row += 1
        entry = values_by_code.get(m["code"])
        label = m["name"] + (" (consolidated)" if entry and entry.is_consolidated else "")
        ws.cell(row=row, column=1, value=label)
        ws.cell(row=row, column=2, value=entry.value if entry else "")
        ws.cell(row=row, column=3, value=m.get("unit", ""))
        row += 1

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf
